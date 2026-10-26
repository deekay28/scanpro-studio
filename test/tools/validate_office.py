#!/usr/bin/env python3
"""Opens Office files with independent libraries and prints a JSON summary.

Usage: validate_office.py <file>...  (docx, xlsx, pptx)
Exit code 1 if any file fails to open. Used by the Dart tests.
"""
import json
import sys


def docx_summary(path):
    import docx
    d = docx.Document(path)
    paras = [p for p in d.paragraphs]
    return {
        "paragraphs": [p.text for p in paras if p.text.strip()],
        "styles": [p.style.name for p in paras if p.text.strip()],
        "bold_runs": [r.text for p in paras for r in p.runs if r.bold],
        "tables": [[[c.text for c in row.cells] for row in t.rows] for t in d.tables],
        "images": len(d.inline_shapes) + sum(1 for r in d.part.rels.values() if "image" in r.reltype),
        "sections": len(d.sections),
        "header": "\n".join(p.text for p in d.sections[0].header.paragraphs),
        "footer": "\n".join(p.text for p in d.sections[0].footer.paragraphs),
    }


def xlsx_summary(path):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    out = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            rows.append([None if v is None else (v.isoformat() if hasattr(v, "isoformat") else v) for v in row])
        out[ws.title] = rows
    return {"sheets": out}


def pptx_summary(path):
    from pptx import Presentation
    p = Presentation(path)
    slides = []
    for s in p.slides:
        info = {"texts": [], "tables": [], "pictures": 0, "charts": []}
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                info["texts"].append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                info["tables"].append([[c.text for c in r.cells] for r in sh.table.rows])
            if sh.shape_type == 13:
                info["pictures"] += 1
            if getattr(sh, "has_chart", False) and sh.has_chart:
                plot = sh.chart.plots[0]
                info["charts"].append({"categories": [str(c) for c in plot.categories], "series": [list(se.values) for se in plot.series]})
        slides.append(info)
    return {"width": p.slide_width, "height": p.slide_height, "slides": slides}


def main():
    result = {}
    ok = True
    for path in sys.argv[1:]:
        try:
            if path.endswith(".docx"):
                result[path] = docx_summary(path)
            elif path.endswith(".xlsx"):
                result[path] = xlsx_summary(path)
            elif path.endswith(".pptx"):
                result[path] = pptx_summary(path)
        except Exception as e:  # noqa: BLE001
            ok = False
            result[path] = {"error": repr(e)}
    print(json.dumps(result, default=str))
    sys.exit(0 if ok else 1)


main()
