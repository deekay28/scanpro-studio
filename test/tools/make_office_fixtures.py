#!/usr/bin/env python3
"""Creates Office test fixtures with python-docx/openpyxl/python-pptx
(a different producer than the app) in the directory given as argv[1]."""
import datetime
import sys

out = sys.argv[1]

# DOCX: headings, formatted runs, list, table, page break.
import docx
from docx.shared import Pt, RGBColor
d = docx.Document()
d.add_heading("Project Charter", level=0)
d.add_heading("Background", level=1)
p = d.add_paragraph("The project starts in ")
r = p.add_run("March 2025")
r.bold = True
p.add_run(" and ends in December.")
r2 = d.add_paragraph().add_run("Red italic note")
r2.italic = True
r2.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
d.add_paragraph("Scope item one", style="List Bullet")
d.add_paragraph("Scope item two", style="List Bullet")
t = d.add_table(rows=3, cols=3)
for i, row in enumerate([["Phase", "Start", "Budget"], ["Design", "2025-03-01", "12000"], ["Build", "2025-05-01", "48000"]]):
    for j, v in enumerate(row):
        t.cell(i, j).text = v
d.add_page_break()
d.add_paragraph("Second page content")
d.sections[0].header.paragraphs[0].text = "ACME Confidential"
d.save(f"{out}/charter.docx")

# XLSX: header, numbers, dates, two sheets.
import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sales"
ws.append(["Region", "Q1", "Q2", "Q3"])
ws.append(["North", 120, 135, 150])
ws.append(["South", 90, 95, 110])
ws.append(["West", 70, 88, 91])
ws2 = wb.create_sheet("Log")
ws2.append(["Date", "Event"])
ws2.append([datetime.date(2025, 1, 15), "Kick-off"])
wb.save(f"{out}/sales.xlsx")

# PPTX: title, bullets, table, chart with embedded data.
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Quarterly Review"
body = s.placeholders[1].text_frame
body.text = "Revenue: 1,250,000"
para = body.add_paragraph()
para.text = "Margin improved"
para.level = 1
s2 = prs.slides.add_slide(prs.slide_layouts[5])
s2.shapes.title.text = "Numbers"
rows, cols = 3, 3
tbl = s2.shapes.add_table(rows, cols, Inches(0.5), Inches(1.5), Inches(4), Inches(1.2)).table
for i, row in enumerate([["Item", "Qty", "Price"], ["Pens", "40", "2.50"], ["Paper", "10", "4.75"]]):
    for j, v in enumerate(row):
        tbl.cell(i, j).text = v
cd = CategoryChartData()
cd.categories = ["Q1", "Q2", "Q3"]
cd.add_series("North", (120, 135, 150))
cd.add_series("South", (90, 95, 110))
s2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1.5), Inches(4.5), Inches(3), cd)
prs.save(f"{out}/review.pptx")
print("ok")
